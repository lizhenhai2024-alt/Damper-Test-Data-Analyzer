from pathlib import Path
import os

import numpy as np
import pandas as pd
import pytest

from cdc_analyzer.dynamic_analysis import TIME, CURRENT, LOAD, DISP, HysteresisConfig, HysteresisStandard
from cdc_analyzer.parser import DataSet
from cdc_analyzer.dynamic_gui_v085 import discover_hysteresis_dat_files
from cdc_analyzer.hysteresis_v085 import analyze_hysteresis_v085, combine_hysteresis_datasets, speed_groups


def multi_speed_data():
    blocks = []
    for speed in (0.1, 0.3, 0.6, 1.0):
        for current, down in ((0.3, False), (0.8, False), (1.6, False), (0.8, True), (0.3, True)):
            block_id = len(blocks) + 1
            t = np.linspace(0, 7, 2801)
            amplitude = speed * 1000 / (2 * np.pi)
            x = -amplitude * np.cos(2 * np.pi * t)
            magnitude = 1000 * speed * (1 + current) + (20 if down else 0)
            force = np.where(np.sin(2 * np.pi * t) >= 0, magnitude, -magnitude * 1.2)
            blocks.append(pd.DataFrame({TIME: t + block_id * 10, DISP: x, CURRENT: current, LOAD: force, "Block ID": block_id}))
    return DataSet(pd.concat(blocks, ignore_index=True), Path("multi_speed.dat"), "synthetic")


@pytest.mark.parametrize("standard", [HysteresisStandard.BMW, HysteresisStandard.AUDI])
def test_hysteresis_separates_speeds_and_does_not_average_them(standard):
    result = analyze_hysteresis_v085(multi_speed_data(), HysteresisConfig(standard=standard))
    assert result.runs["Speed Group m/s"].nunique() == 4
    assert result.summary["Speed Group m/s"].nunique() == 4
    rebound = result.summary[result.summary["Direction"] == "Rebound"]
    if standard == HysteresisStandard.BMW:
        rebound = rebound[rebound["Current A"] == 0.8]
    assert len(rebound) == 4
    assert rebound["Hysteresis N"].to_numpy() == pytest.approx([20] * 4, abs=1)


def test_speed_groups_do_not_chain_distinct_conditions():
    assert len(np.unique(speed_groups([0.1, 0.102, 0.104, 0.3]))) == 3


def test_combines_separate_speed_files_without_block_id_collisions():
    combined = multi_speed_data()
    files = []
    for index, block_ids in enumerate(np.array_split(combined.data["Block ID"].unique(), 4), start=1):
        frame = combined.data[combined.data["Block ID"].isin(block_ids)].copy()
        frame["Block ID"] -= int(frame["Block ID"].min()) - 1
        files.append(DataSet(frame, Path(f"speed_{index}.dat"), "synthetic"))
    merged = combine_hysteresis_datasets(files)
    assert merged.data["Block ID"].nunique() == 20
    assert merged.metadata["source_file_count"] == 4
    assert merged.data["Source File"].nunique() == 4
    result = analyze_hysteresis_v085(merged, HysteresisConfig(standard=HysteresisStandard.BMW))
    assert result.runs["Speed Group m/s"].nunique() == 4


def test_hysteresis_folder_scan_finds_dat_files_recursively(tmp_path):
    nested = tmp_path / "speed" / "compression"
    nested.mkdir(parents=True)
    first = tmp_path / "root.dat"
    second = nested / "specimen.DAT"
    ignored = nested / "notes.csv"
    for path in (first, second, ignored):
        path.write_text("test", encoding="utf-8")

    assert discover_hysteresis_dat_files(tmp_path) == [first, second]


def test_v085_gui_response_and_hysteresis(tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6 import QtCore, QtWidgets
    import pyqtgraph as pg
    from cdc_analyzer.gui_release_v085 import _build_release_gui_classes_v085
    from test_gui_v083 import _fake_response_result
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    window = _build_release_gui_classes_v085()()
    window.resize(1280, 800)
    pages = window.dynamic_pages
    assert pages.response_force_start_fraction.value() == pytest.approx(1.0)
    assert pages.response_show_f1.isChecked()
    assert pages.response_show_f63.isChecked()
    pages.response_result = _fake_response_result()
    pages._rebuild_response_event_combo()
    window.tabs.setCurrentWidget(pages.response_page)
    pages.refresh_response_plot()
    window.show()
    for _ in range(3):
        app.processEvents()
    for manager in pages.response_annotations:
        manager.update()
        for item, x, y, level in manager.labels:
            assert not item.textItem.font().bold()
            assert item.textItem.font().pointSizeF() == manager.plot.getAxis("left").label.font().pointSizeF()
            assert item.fill.style().value == 0
            assert item.border.style().value == 0
        rects = manager.text_rects
        for i, rect in enumerate(rects):
            assert all(not rect.intersects(other) for other in rects[i + 1:])
        level_rects = [
            rect for (_, _, _, level), rect in zip(manager.labels, rects) if level
        ]
        if len(level_rects) > 1:
            assert max(rect.left() for rect in level_rects) - min(
                rect.left() for rect in level_rects
            ) < 1
            level_geometry = [
                (
                    manager.plot.vb.mapViewToScene(QtCore.QPointF(x, y)).y(),
                    rect.center().y(),
                )
                for (_, x, y, level), rect in zip(manager.labels, rects)
                if level
            ]
            by_guide = sorted(level_geometry)
            # Labels must preserve the screen order of their own guide lines.
            assert [label_y for _, label_y in by_guide] == sorted(
                label_y for _, label_y in by_guide
            )
        assert len(manager.lines) == len(manager.guides)
        for (_, x, y, level), rect in zip(manager.labels, rects):
            if not level:
                point = manager.plot.vb.mapViewToScene(QtCore.QPointF(x, y))
                assert not rect.contains(point)
        dots = [item for item in manager.plot.items if isinstance(item, pg.ScatterPlotItem)]
        assert len(dots) == 1
        for point in dots[0].points():
            signal = manager.plot.listDataItems()[0]
            assert point.pos().y() == pytest.approx(np.interp(point.pos().x(), signal.xData, signal.yData))
    force_annotations = pages.response_annotations[1]
    response_placements = [
        force_annotations.placements[item]
        for item, _, _, level in force_annotations.labels
        if not level and item.textItem.toPlainText().startswith(("t₁%", "t₆₃%", "t₉₀%"))
    ]
    assert [placement.split("-", 1)[0] for placement in response_placements] == [
        "above", "below", "above"
    ]
    pages.hysteresis_result = analyze_hysteresis_v085(multi_speed_data(), HysteresisConfig(standard=HysteresisStandard.BMW))
    pages._fill_table(pages.hysteresis_run_table, pages.hysteresis_result.runs)
    force_column = pages.hysteresis_result.runs.columns.get_loc("Force N")
    expected_force = f"{pages.hysteresis_result.runs.iloc[0]['Force N']:.0f}"
    assert pages.hysteresis_run_table.item(0, force_column).text() == expected_force
    pages._rebuild_hysteresis_views()
    assert pages.hysteresis_multi_file_button.isVisibleTo(pages.hysteresis_page)
    assert pages.hysteresis_folder_button.isVisibleTo(pages.hysteresis_page)
    assert pages.hysteresis_view_combo.count() == 1 + 1 + 4 + 3
    plot = pages.hysteresis_plot_area.getItem(0, 0)
    assert plot.getAxis("left").label.toPlainText().strip() == "压缩<--阻尼力(N)-->复原"
    assert all(
        curve.opts["pen"].style() == QtCore.Qt.PenStyle.SolidLine
        for curve in plot.listDataItems()
    )
    assert all(curve.opts["symbol"] is None for curve in plot.listDataItems())
    assert all(curve.opts["connect"] == "all" for curve in plot.listDataItems())
    assert all(curve.opts["antialias"] is True for curve in plot.listDataItems())
    ys = np.concatenate([p.yData for p in plot.listDataItems()])
    assert ys.min() < 0 < ys.max()
    for i in range(pages.hysteresis_view_combo.count()):
        pages.hysteresis_view_combo.setCurrentIndex(i)
        selected_mode = pages.hysteresis_view_combo.currentData()[0]
        selected_plot = pages.hysteresis_plot_area.getItem(0, 0)
        if selected_mode == "hysteresis_bar":
            bars = [item for item in selected_plot.items if isinstance(item, pg.BarGraphItem)]
            assert bars
            assert all(np.all(np.asarray(item.opts["height"]) >= 0) for item in bars)
            labels = [item for item in selected_plot.items if isinstance(item, pg.TextItem)]
            assert labels
            assert all(item.toPlainText().isdigit() for item in labels)
            continue
        curves = selected_plot.listDataItems()
        assert curves
        assert all(curve.opts["pen"].style() == QtCore.Qt.PenStyle.SolidLine for curve in curves)
        assert all(curve.opts["symbol"] is None for curve in curves)
        assert all(curve.opts["connect"] == "all" for curve in curves)
    window.tabs.setCurrentIndex(0)
    pages.hysteresis_view_combo.setCurrentIndex(0)
    pages._prepare_hysteresis_plot_for_export()
    plot = pages.hysteresis_plot_area.getItem(0, 0)
    ys = np.concatenate([p.yData for p in plot.listDataItems()])
    y_range = plot.viewRange()[1]
    assert y_range[0] < np.nanmin(ys) < 0 < np.nanmax(ys) < y_range[1]
    from cdc_analyzer.dynamic_export import export_hysteresis_xlsx
    from openpyxl import load_workbook
    workbook_path = export_hysteresis_xlsx(pages.hysteresis_result, tmp_path / "grouped.xlsx")
    exported_book = load_workbook(workbook_path)
    run_sheet = exported_book["Run Detail"]
    force_headers = {cell.value: cell.column for cell in run_sheet[1]}
    for column in ("Force N", "Abs Force N"):
        assert run_sheet.cell(2, force_headers[column]).number_format == "0"
    summary_sheet = exported_book["Hysteresis Summary"]
    summary_headers = {cell.value: cell.column for cell in summary_sheet[1]}
    for column in ("Up Force N", "Down Force N", "Reference Damping Force N", "Hysteresis N"):
        assert summary_sheet.cell(2, summary_headers[column]).number_format == "0"
    exported_book.close()
    previous = pages.hysteresis_view_combo.currentIndex()
    pages._append_hysteresis_plot(workbook_path)
    assert pages.hysteresis_view_combo.currentIndex() == previous
    book = load_workbook(workbook_path)
    assert len(book["Hysteresis Plot"]._images) == 9
    book.close()
    window.close()
    app.processEvents()


def test_current_packaged_gui_is_v085():
    root = Path(__file__).resolve().parents[1]
    assert "gui_release_v085" in (root / "launcher.py").read_text()
    assert "gui_release_v085:main" in (root / "pyproject.toml").read_text()
    workflow = (root / ".github" / "workflows" / "build-windows.yml").read_text()
    assert "APP_VERSION: V0.8.20" in workflow
    assert 'Damper_Test_Data_Analyzer_$env:APP_VERSION' in workflow
    assert "Damper_Test_Data_Analyzer_${{ env.APP_VERSION }}.exe" in workflow


def test_dual_axis_current_response_plot_marks_both_signals():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6 import QtCore, QtWidgets
    from cdc_analyzer.gui_release_v085 import _build_release_gui_classes_v085
    from test_gui_v083 import _fake_response_result

    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    window = _build_release_gui_classes_v085()()
    pages = window.dynamic_pages
    result = _fake_response_result()
    result.events["I10 Crossing Time s"] = 1.110
    result.events["I90 Crossing Time s"] = 1.117
    result.events["Current Response I10-I90 ms"] = 7.0
    result.events["F90 Crossing Time s"] = 1.130
    result.events["Damping Response I10-F90 ms"] = 20.0
    pages.response_result = result
    pages._rebuild_response_event_combo()

    assert pages.response_plot_mode.count() == 2
    assert pages.response_plot_mode.currentData() == "stacked"
    pages.response_plot_mode.setCurrentIndex(
        pages.response_plot_mode.findData("dual_axis_i10_f90")
    )
    pages.refresh_response_plot()

    items = pages.response_dual_items
    assert len(items["guides"]) == 2
    assert all(
        guide.pen.style() == QtCore.Qt.PenStyle.DashLine
        for guide in items["guides"]
    )
    assert len(items["current_points"].points()) == 2
    assert len(items["force_points"].points()) == 2
    assert [point.pos().x() for point in items["current_points"].points()] == pytest.approx(
        [1.110, 1.130]
    )
    expected_force = np.interp(
        [1.110, 1.130],
        result.processed[TIME],
        result.processed[LOAD] / 1000.0,
    )
    assert [point.pos().y() for point in items["force_points"].points()] == pytest.approx(
        expected_force
    )
    assert items["current_curve"].opts["pen"].color().name() == "#1565c0"
    assert items["force_curve"].opts["pen"].color().name() == "#c62828"
    plot = items["plot"]
    assert plot.getAxis("left").textPen().color().name() == "#1565c0"
    assert plot.getAxis("right").textPen().color().name() == "#c62828"
    texts = [label.toPlainText() for label in items["labels"]]
    assert "I₁₀%" in texts
    assert any(text.startswith("I(F₉₀%) =") for text in texts)
    assert any(text.startswith("F(I₁₀%) =") for text in texts)
    assert any(text.startswith("F₉₀% =") for text in texts)
    assert "t₉₀% = t(F₉₀%) − t(I₁₀%) = 20.00 ms" in texts
    assert all(not label.textItem.font().bold() for label in items["labels"])
    assert all(label.fill.style().value == 0 for label in items["labels"])

    window.close()
    app.processEvents()


def test_response_threshold_labels_follow_settings_and_visibility():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6 import QtWidgets
    from cdc_analyzer.gui_release_v085 import _build_release_gui_classes_v085
    from test_gui_v083 import _fake_response_result

    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    window = _build_release_gui_classes_v085()()
    pages = window.dynamic_pages
    result = _fake_response_result()
    result.events["Trigger Fraction"] = 0.25
    result.events["Force Start Fraction"] = 0.05
    pages.response_result = result
    pages.response_force_start_fraction.setValue(5.0)
    pages._rebuild_response_event_combo()
    pages.refresh_response_plot()

    current_labels = [
        (item.textItem.toPlainText(), x, level)
        for item, x, _y, level in pages.response_annotations[0].labels
    ]
    assert ("I₂₅%", float(result.events.iloc[0]["t0 s"]), False) in current_labels
    assert not any(text == "I₂₅%" and level for text, _x, level in current_labels)

    force_texts = [
        item.textItem.toPlainText()
        for item, _x, _y, _level in pages.response_annotations[1].labels
    ]
    assert "F₅%" in force_texts
    assert any(text.startswith("t₅% =") for text in force_texts)
    assert "F₆₃%" in force_texts

    pages.response_show_f1.setChecked(False)
    pages.response_show_f63.setChecked(False)
    force_texts = [
        item.textItem.toPlainText()
        for item, _x, _y, _level in pages.response_annotations[1].labels
    ]
    assert "F₅%" not in force_texts
    assert "F₆₃%" not in force_texts
    assert not any(text.startswith(("t₅% =", "t₆₃% =")) for text in force_texts)
    assert "F₉₀%" in force_texts
    assert "F₁₀₀%" in force_texts

    window.close()
    app.processEvents()


def test_main_evaluation_method_excludes_audi_and_defaults_to_window_mean():
    from PySide6 import QtWidgets
    from cdc_analyzer.analysis import EvaluationProfile
    from cdc_analyzer.gui_release_v085 import _build_release_gui_classes_v085

    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    window = _build_release_gui_classes_v085()()
    methods = [window.profile.itemData(index) for index in range(window.profile.count())]

    assert methods == [
        EvaluationProfile.WINDOW_MEAN.value,
        EvaluationProfile.ZERO_CROSSING.value,
    ]
    assert window.profile.currentData() == EvaluationProfile.WINDOW_MEAN.value
    assert window.window_basis.currentData() == "total_stroke"
    assert window.window_basis.count() == 1
    assert window.window_basis.isHidden()
    assert window.eval_form.labelForField(window.window_basis).isHidden()
    assert window._config().window_basis == "total_stroke"

    window.language_combo.setCurrentIndex(window.language_combo.findData("en_US"))
    app.processEvents()
    assert [window.profile.itemData(index) for index in range(window.profile.count())] == methods
    assert window.profile.currentData() == EvaluationProfile.WINDOW_MEAN.value

    window.close()
    app.processEvents()


@pytest.mark.parametrize(
    ("levels", "force_curve", "f100_below_f90"),
    [
        ((-1310, -2050, -2850, -3000), (-1300, -3000), True),
        ((900, 1800, 2800, 3000), (900, 3000), False),
    ],
)
def test_force_level_labels_preserve_compression_and_rebound_order(
    levels, force_curve, f100_below_f90
):
    from PySide6 import QtWidgets
    from cdc_analyzer.gui_release_v085 import _build_release_gui_classes_v085
    from test_gui_v083 import _fake_response_result

    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    window = _build_release_gui_classes_v085()()
    window.resize(1280, 800)
    pages = window.dynamic_pages
    result = _fake_response_result()
    result.events.loc[0, ["F1 N", "F63 N", "F90 N", "F100 N"]] = levels
    result.processed[LOAD] = np.linspace(*force_curve, len(result.processed))
    pages.response_result = result
    pages._rebuild_response_event_combo()
    pages.refresh_response_plot()
    window.show()
    for _ in range(3):
        app.processEvents()

    manager = pages.response_annotations[1]
    manager.update()
    label_y = {
        item.textItem.toPlainText(): rect.center().y()
        for (item, _x, _y, level), rect in zip(manager.labels, manager.text_rects)
        if level
    }
    if f100_below_f90:
        assert label_y["F₁₀₀%"] > label_y["F₉₀%"]
    else:
        assert label_y["F₁₀₀%"] < label_y["F₉₀%"]

    window.close()
    app.processEvents()


@pytest.mark.parametrize("scale", ["1", "1.5"])
def test_display_scale_keeps_controls_and_annotations_readable(scale):
    import subprocess
    import sys
    root = Path(__file__).resolve().parents[1]
    env = dict(os.environ, QT_QPA_PLATFORM="offscreen", QT_SCALE_FACTOR=scale,
               PYTHONPATH=str(root / "src") + os.pathsep + str(root / "tests"))
    completed = subprocess.run([sys.executable, str(root / "tests" / "probe_display_v085.py")],
                               env=env, capture_output=True, text=True, timeout=120)
    assert completed.returncode == 0, completed.stdout + completed.stderr
