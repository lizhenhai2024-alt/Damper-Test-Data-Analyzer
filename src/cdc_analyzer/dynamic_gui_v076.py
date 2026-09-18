from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

from .dynamic_export import export_hysteresis_xlsx, export_response_xlsx
from .dynamic_gui_v075 import DynamicPagesController as _V075DynamicPagesController
from .image_export import DEFAULT_PNG_EXPORT_PPI, export_plot_widget_png


class DynamicPagesController(_V075DynamicPagesController):
    """V0.7.6 dynamic pages.

    Response and hysteresis now consume the shared data source loaded from the
    main left-side Data File panel. Module-local Open/Export buttons remain as
    compatibility objects but are hidden from the user-facing layout.
    """

    def _build_response_page(self):
        super()._build_response_page()
        self.response_open_button.hide()
        self.response_export_button.hide()

    def _build_hysteresis_page(self):
        super()._build_hysteresis_page()
        self.hysteresis_open_button.hide()
        self.hysteresis_export_button.hide()

    def set_shared_dataset(self, dataset, path: str | Path | None = None) -> None:
        source_path = Path(path) if path is not None else Path(dataset.source_path)
        self.response_dataset = dataset
        self.response_path = source_path
        self.hysteresis_dataset = dataset
        self.hysteresis_path = source_path
        self.response_result = None
        self.hysteresis_result = None
        self._update_shared_source_labels()
        self.response_status.setText(
            self._text("已使用左侧公共数据源；点击“分析响应”开始计算。", "Shared data source loaded; click Analyze Response.")
        )
        self.hysteresis_status.setText(
            self._text("已使用左侧公共数据源；点击“分析迟滞”开始计算。", "Shared data source loaded; click Analyze Hysteresis.")
        )

    def _update_shared_source_labels(self) -> None:
        if self.response_path is not None:
            self.response_file_label.setText(
                self._text(f"公共数据源：{self.response_path}", f"Shared source: {self.response_path}")
            )
        else:
            self.response_file_label.setText(
                self._text("公共数据源：未加载", "Shared source: not loaded")
            )
        if self.hysteresis_path is not None:
            self.hysteresis_file_label.setText(
                self._text(f"公共数据源：{self.hysteresis_path}", f"Shared source: {self.hysteresis_path}")
            )
        else:
            self.hysteresis_file_label.setText(
                self._text("公共数据源：未加载", "Shared source: not loaded")
            )

    def apply_language(self, language: str):
        super().apply_language(language)
        if hasattr(self, "response_open_button"):
            self.response_open_button.hide()
            self.response_export_button.hide()
            self.hysteresis_open_button.hide()
            self.hysteresis_export_button.hide()
            self._update_shared_source_labels()

    def _export_plot_widget_png(self, plot_widget, path: str | Path, scale: float = 3.0) -> Path:
        ppi = getattr(self.window, "png_export_ppi", DEFAULT_PNG_EXPORT_PPI)
        return export_plot_widget_png(plot_widget, path, scale, ppi)

    def _png_export_ppi(self) -> int:
        return int(getattr(self.window, "png_export_ppi", DEFAULT_PNG_EXPORT_PPI))

    def export_response_png(self) -> None:
        if self.response_result is None:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("导出 PNG", "Export PNG"),
                self._text("请先完成响应时间分析。", "Analyze response time first."),
            )
            return
        default = (
            self.response_path.stem + "_response.png"
            if self.response_path is not None
            else "response.png"
        )
        path, _ = self.QtWidgets.QFileDialog.getSaveFileName(
            self.window,
            self._text("导出响应时间图片", "Export response-time plot"),
            default,
            "PNG (*.png)",
        )
        if not path:
            return
        try:
            if hasattr(self, "response_view_tabs"):
                self.response_view_tabs.setCurrentWidget(self.response_graph_page)
            self.refresh_response_plot()
            self.QtWidgets.QApplication.processEvents()
            out = self._export_plot_widget_png(self.response_plot_area, path, 3.0)
            self.window.statusBar().showMessage(
                self._text(
                    f"已导出 {self._png_export_ppi()} PPI PNG：{out}",
                    f"{self._png_export_ppi()}-PPI PNG exported: {out}",
                )
            )
        except Exception as exc:
            self.QtWidgets.QMessageBox.critical(
                self.window,
                self._text("导出错误", "Export error"),
                str(exc),
            )

    def export_hysteresis_png(self) -> None:
        if self.hysteresis_result is None:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("导出 PNG", "Export PNG"),
                self._text("请先完成迟滞分析。", "Analyze hysteresis first."),
            )
            return
        default = (
            self.hysteresis_path.stem + "_hysteresis.png"
            if self.hysteresis_path is not None
            else "hysteresis.png"
        )
        path, _ = self.QtWidgets.QFileDialog.getSaveFileName(
            self.window,
            self._text("导出迟滞图片", "Export hysteresis plot"),
            default,
            "PNG (*.png)",
        )
        if not path:
            return
        try:
            self.refresh_hysteresis_plot()
            self.QtWidgets.QApplication.processEvents()
            out = self._export_plot_widget_png(self.hysteresis_plot_area, path, 3.0)
            self.window.statusBar().showMessage(
                self._text(
                    f"已导出 {self._png_export_ppi()} PPI PNG：{out}",
                    f"{self._png_export_ppi()}-PPI PNG exported: {out}",
                )
            )
        except Exception as exc:
            self.QtWidgets.QMessageBox.critical(
                self.window,
                self._text("导出错误", "Export error"),
                str(exc),
            )

    def _append_response_plots(self, workbook_path: Path) -> None:
        if self.response_result is None or self.response_result.events.empty:
            return
        original_index = self.response_event_combo.currentIndex()
        workbook = load_workbook(workbook_path)
        if "Response Plots" in workbook.sheetnames:
            del workbook["Response Plots"]
        sheet = workbook.create_sheet("Response Plots")
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 24

        with TemporaryDirectory(prefix="damper_response_plots_") as tmp:
            tmp_dir = Path(tmp)
            anchor_row = 1
            for index, row in self.response_result.events.reset_index(drop=True).iterrows():
                event_id = int(row["Event ID"])
                combo_index = self.response_event_combo.findData(event_id)
                if combo_index >= 0:
                    self.response_event_combo.setCurrentIndex(combo_index)
                self.refresh_response_plot()
                self.QtWidgets.QApplication.processEvents()
                image_path = tmp_dir / f"event_{event_id:03d}.png"
                self._export_plot_widget_png(self.response_plot_area, image_path, 2.5)

                stage = str(row.get("Stage", ""))
                direction = str(row.get("Direction", ""))
                current_transition = str(row.get("Current Transition", ""))
                target = row.get("Target Velocity m/s", "")
                sheet.cell(anchor_row, 1).value = (
                    f"Event {event_id} | {stage} | {direction} | {current_transition} | Target {target} m/s"
                )
                image = XLImage(str(image_path))
                if image.width:
                    ratio = min(1.0, 1180.0 / float(image.width))
                    image.width = int(image.width * ratio)
                    image.height = int(image.height * ratio)
                sheet.add_image(image, f"A{anchor_row + 1}")
                rows_used = max(32, int(image.height / 20) + 5)
                anchor_row += rows_used

            # openpyxl reads image files during save; keep the temporary PNGs alive.
            workbook.save(workbook_path)

        if original_index >= 0 and original_index < self.response_event_combo.count():
            self.response_event_combo.setCurrentIndex(original_index)
            self.refresh_response_plot()

    def _append_hysteresis_plot(self, workbook_path: Path) -> None:
        if self.hysteresis_result is None:
            return
        workbook = load_workbook(workbook_path)
        if "Hysteresis Plot" in workbook.sheetnames:
            del workbook["Hysteresis Plot"]
        sheet = workbook.create_sheet("Hysteresis Plot")
        sheet.sheet_view.showGridLines = False
        with TemporaryDirectory(prefix="damper_hysteresis_plot_") as tmp:
            image_path = Path(tmp) / "hysteresis.png"
            self.refresh_hysteresis_plot()
            self.QtWidgets.QApplication.processEvents()
            self._export_plot_widget_png(self.hysteresis_plot_area, image_path, 2.5)
            image = XLImage(str(image_path))
            if image.width:
                ratio = min(1.0, 1180.0 / float(image.width))
                image.width = int(image.width * ratio)
                image.height = int(image.height * ratio)
            sheet.add_image(image, "A1")
            workbook.save(workbook_path)

    def export_response(self):
        if self.response_result is None:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("导出 Excel", "Export Excel"),
                self._text("请先完成响应时间分析。", "Analyze response time first."),
            )
            return
        default = (
            self.response_path.stem + "_response.xlsx"
            if self.response_path is not None
            else "response.xlsx"
        )
        path, _ = self.QtWidgets.QFileDialog.getSaveFileName(
            self.window,
            self._text("导出响应时间结果", "Export response-time result"),
            default,
            "Excel (*.xlsx)",
        )
        if not path:
            return
        try:
            out = export_response_xlsx(self.response_result, path)
            self._append_response_plots(out)
            self.window.statusBar().showMessage(
                self._text(f"已导出响应 Excel（全部阶段与图形）：{out}", f"Response Excel exported with all event plots: {out}")
            )
        except Exception as exc:
            self.QtWidgets.QMessageBox.critical(
                self.window,
                self._text("导出错误", "Export error"),
                str(exc),
            )

    def export_hysteresis(self):
        if self.hysteresis_result is None:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("导出 Excel", "Export Excel"),
                self._text("请先完成迟滞分析。", "Analyze hysteresis first."),
            )
            return
        default = (
            self.hysteresis_path.stem + "_hysteresis.xlsx"
            if self.hysteresis_path is not None
            else "hysteresis.xlsx"
        )
        path, _ = self.QtWidgets.QFileDialog.getSaveFileName(
            self.window,
            self._text("导出迟滞结果", "Export hysteresis result"),
            default,
            "Excel (*.xlsx)",
        )
        if not path:
            return
        try:
            out = export_hysteresis_xlsx(self.hysteresis_result, path)
            self._append_hysteresis_plot(out)
            self.window.statusBar().showMessage(
                self._text(f"已导出迟滞 Excel（含图形）：{out}", f"Hysteresis Excel exported with plot: {out}")
            )
        except Exception as exc:
            self.QtWidgets.QMessageBox.critical(
                self.window,
                self._text("导出错误", "Export error"),
                str(exc),
            )
