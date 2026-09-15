from __future__ import annotations

import io
import re
import struct
import zlib
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator
from xml.etree import ElementTree

import numpy as np
import pandas as pd

from .analysis import AnalyzerConfig, CURRENT, DISP, EvaluationProfile, LOAD, TIME, detect_complete_cycles


@dataclass(slots=True)
class ImportedMapFile:
    path: Path
    current_a: float
    format: str
    run_count: int
    speeds_mps: tuple[float, ...]
    status: str = "OK"
    error: str = ""


@dataclass(slots=True)
class MapAnalysisResult:
    files: pd.DataFrame
    run_detail: pd.DataFrame
    current_force_linearity: pd.DataFrame
    spread_amplification: pd.DataFrame
    settings: dict[str, object]


def discover_map_files(folder: str | Path) -> list[Path]:
    root = Path(folder)
    return sorted(
        (p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in {".pvp", ".dctw"}),
        key=lambda p: str(p).casefold(),
    )


def current_from_filename(path: str | Path) -> float:
    stem = Path(path).stem.replace(",", ".")
    matches = re.findall(r"(?<![\d.])(\d+(?:\.\d+)?)\s*[aA](?=$|[-_\s])", stem)
    if matches:
        return float(matches[-1])
    if re.fullmatch(r"\s*\d+(?:\.\d+)?\s*", stem):
        return float(stem)
    raise ValueError(f"无法从文件名识别电流，请在列表中填写电流 A: {Path(path).name}")


class _BondReader:
    """Minimal Bond Compact Binary v1 reader used by CTW .dctw files."""

    def __init__(self, data: bytes):
        self.stream = io.BytesIO(data)

    def tell(self) -> int:
        return self.stream.tell()

    def read(self, size: int) -> bytes:
        value = self.stream.read(size)
        if len(value) != size:
            raise EOFError
        return value

    def u8(self) -> int:
        return self.read(1)[0]

    def varuint(self) -> int:
        value = shift = 0
        while True:
            byte = self.u8()
            value |= (byte & 0x7F) << shift
            if byte < 0x80:
                return value
            shift += 7
            if shift > 70:
                raise ValueError("Invalid Bond variable integer")

    def value(self, data_type: int):
        if data_type in (2, 3, 14):
            return self.u8()
        if data_type in (4, 5, 6):
            return self.varuint()
        if data_type in (15, 16, 17):
            value = self.varuint()
            return (value >> 1) ^ -(value & 1)
        if data_type == 7:
            return struct.unpack("<f", self.read(4))[0]
        if data_type == 8:
            return struct.unpack("<d", self.read(8))[0]
        if data_type == 9:
            return self.read(self.varuint()).decode("utf-8", errors="replace")
        if data_type == 18:
            return self.read(self.varuint() * 2).decode("utf-16le", errors="replace")
        if data_type == 10:
            return self.structure()
        if data_type in (11, 12):
            element_type = self.u8() & 0x1F
            return [self.value(element_type) for _ in range(self.varuint())]
        if data_type == 13:
            key_type, value_type, count = self.u8(), self.u8(), self.varuint()
            return [(self.value(key_type), self.value(value_type)) for _ in range(count)]
        raise ValueError(f"Unsupported Bond data type {data_type} at {self.tell()}")

    def structure(self) -> list[tuple[int | str, int | None, object]]:
        fields: list[tuple[int | str, int | None, object]] = []
        while True:
            header = self.u8()
            data_type, short_id = header & 0x1F, header >> 5
            if data_type == 0:
                return fields
            if data_type == 1:
                fields.append(("BASE", None, None))
                continue
            field_id = short_id if short_id <= 5 else self.u8() if short_id == 6 else struct.unpack("<H", self.read(2))[0]
            fields.append((field_id, data_type, self.value(data_type)))


def _layers(fields) -> list[list[tuple[int, int, object]]]:
    result, current = [], []
    for field in fields:
        if field[0] == "BASE":
            result.append(current)
            current = []
        else:
            current.append(field)
    result.append(current)
    return result


def _field(layer, field_id: int, default=None):
    for identifier, _data_type, value in layer:
        if identifier == field_id:
            return value
    return default


def _all_strings(value) -> Iterator[str]:
    if isinstance(value, str):
        yield value
    elif isinstance(value, list):
        for item in value:
            if isinstance(item, tuple) and len(item) == 3:
                yield from _all_strings(item[2])
            else:
                yield from _all_strings(item)
    elif isinstance(value, tuple):
        for item in value:
            yield from _all_strings(item)


def _dctw_records(path: Path):
    raw = path.read_bytes()
    reader = _BondReader(raw)
    header = reader.structure()
    header_end = reader.tell()
    if header_end + 16 > len(raw):
        raise ValueError("CTW file header is incomplete")
    element_count, uncompressed_start = struct.unpack("<Qq", raw[header_end : header_end + 16])
    position = header_end + 16
    prefix = _BondReader(raw[position:uncompressed_start])
    compressed_offset = None
    while prefix.tell() < uncompressed_start - position:
        element_type = prefix.u8()
        fields = prefix.structure()
        if element_type == 39 and bool(_field(_layers(fields)[-1], 0, False)):
            compressed_offset = position + prefix.tell()
            break
        yield element_type, fields
    if compressed_offset is None:
        raise ValueError("CTW compressed data block was not found")
    try:
        expanded = zlib.decompress(raw[compressed_offset:uncompressed_start], -zlib.MAX_WBITS)
    except zlib.error as exc:
        raise ValueError(f"CTW compressed data is invalid: {exc}") from exc
    for block in (expanded, raw[uncompressed_start:]):
        stream = _BondReader(block)
        while stream.tell() < len(block):
            try:
                element_type = stream.u8()
                fields = stream.structure()
            except EOFError:
                break
            if element_type != 39:
                yield element_type, fields
    if element_count == 0:
        raise ValueError("CTW file contains no elements")
    return header


def _extract_dctw_header(path: Path):
    raw = path.read_bytes()
    reader = _BondReader(raw)
    fields = reader.structure()
    xml_candidates = [s.lstrip("\ufeff") for s in _all_strings(fields) if "<" in s]
    xml_text = max(xml_candidates, key=len, default="").rstrip("\x00")
    if not xml_text:
        raise ValueError("CTW machine configuration was not found")
    root = ElementTree.fromstring(xml_text)
    channels = {}
    for channel in root.findall(".//Channels/Channel"):
        try:
            signal_id = int(channel.findtext("Id", ""))
        except ValueError:
            continue
        properties = channel.find("Properties")
        name = properties.findtext("Name/string", "") if properties is not None else ""
        units = properties.findtext("Units/string", "") if properties is not None else ""
        calibration = []
        for entry in channel.findall("Calibration/CalibrationEntry"):
            try:
                calibration.append((float(entry.findtext("IndependentValue", "")), float(entry.findtext("DependantValue", ""))))
            except ValueError:
                continue
        channels[signal_id] = {"name": name, "units": units, "calibration": sorted(calibration)}
    return channels


def _calibrate(values: np.ndarray, points: list[tuple[float, float]]) -> np.ndarray:
    if len(points) < 2:
        return values.astype(float, copy=True)
    xp, fp = (np.asarray(v, dtype=float) for v in zip(*points))
    output = np.interp(values, xp, fp)
    left = values < xp[0]
    right = values > xp[-1]
    output[left] = fp[0] + (values[left] - xp[0]) * (fp[1] - fp[0]) / (xp[1] - xp[0])
    output[right] = fp[-1] + (values[right] - xp[-1]) * (fp[-1] - fp[-2]) / (xp[-1] - xp[-2])
    return output


def _parse_dctw(path: Path, current_a: float):
    channels = _extract_dctw_header(path)
    by_name = {value["name"].casefold(): key for key, value in channels.items()}
    force_id = by_name.get("force", 0)
    displacement_id = by_name.get("displacement", 8)
    velocity_id = by_name.get("velocity", 11)
    force_info = channels.get(force_id, {})
    active = None
    runs = []
    for element_type, fields in _dctw_records(path):
        layers = _layers(fields)
        if element_type == 33:
            strings = list(_all_strings(fields))
            text = next((s for s in strings if "Run Test Speed" in s), "")
            match = re.search(r"at\s*\[\s*([0-9.]+)\s*m/s\s*\]", text, re.I)
            if match:
                active = {"speed": float(match.group(1)), "signals": {force_id: [], displacement_id: [], velocity_id: []}}
        elif element_type == 4 and active is not None:
            derived = layers[-1]
            for signal in _field(derived, 1, []):
                values = {identifier: value for identifier, _type, value in signal if identifier != "BASE"}
                # Bond omits fields holding their schema default; SignalId 0 is
                # therefore absent for the first (force) channel.
                signal_id = int(values.get(0, 0))
                if signal_id in active["signals"]:
                    active["signals"][signal_id].extend(values.get(1, []))
        elif element_type == 2 and active is not None:
            if any("RunSpeed" in s for s in _all_strings(fields)):
                signals = active["signals"]
                count = min((len(signals[key]) for key in signals), default=0)
                if count:
                    force = np.asarray(signals[force_id][:count], dtype=float)
                    force = _calibrate(force, force_info.get("calibration", []))
                    if str(force_info.get("units", "")).casefold() in {"lb", "lbs", "lbf"}:
                        force *= 4.4482216152605
                    displacement = np.asarray(signals[displacement_id][:count], dtype=float)
                    if str(channels.get(displacement_id, {}).get("units", "")).casefold() in {"in", "inch", "inches"}:
                        displacement *= 25.4
                    velocity = np.asarray(signals[velocity_id][:count], dtype=float)
                    if str(channels.get(velocity_id, {}).get("units", "")).casefold() in {"in/s", "ips"}:
                        velocity *= 0.0254
                    runs.append((active["speed"], displacement, force, velocity))
                active = None
    if not runs:
        raise ValueError("CTW file contains no complete RunSpeed data")
    return runs


def _pvp_channel_records(data: bytes, name: bytes):
    records = []
    start = 0
    while True:
        index = data.find(name, start)
        if index < 0:
            return records
        start = index + 1
        if index < 4 or struct.unpack_from("<I", data, index - 4)[0] != len(name):
            continue
        header = index + len(name)
        if header + 16 > len(data):
            continue
        _offset, _unit, reserved, count = struct.unpack_from("<IIII", data, header)
        data_start = header + 16
        if reserved == 0 and 1 <= count <= 20_000_000 and data_start + 4 * count <= len(data):
            values = np.frombuffer(data, dtype="<f4", count=count, offset=data_start).astype(float)
            records.append((index, values))
    return records


def _parse_pvp(path: Path, current_a: float):
    try:
        import olefile
    except ImportError as exc:  # pragma: no cover - dependency is packaged
        raise RuntimeError("Reading PVP files requires the olefile package") from exc
    if not olefile.isOleFile(str(path)):
        raise ValueError("PVP file is not a valid OLE compound document")
    with olefile.OleFileIO(str(path)) as document:
        data = document.openstream("Session").read()
    displacements = _pvp_channel_records(data, b"Displacement")
    forces = _pvp_channel_records(data, b"Force")
    velocities = _pvp_channel_records(data, b"Velocity")
    runs = []
    for index, (position, displacement) in enumerate(displacements):
        end = displacements[index + 1][0] if index + 1 < len(displacements) else len(data)
        force = next((v for p, v in forces if position < p < end), None)
        velocity = next((v for p, v in velocities if position < p < end), None)
        if force is None or velocity is None:
            continue
        count = min(len(displacement), len(force), len(velocity))
        before = data[max(0, position - 5000):position]
        speed_at = before.rfind(b"Test_Speed")
        speed = float(np.nanmax(np.abs(velocity[:count])))
        if speed_at >= 0:
            tail = before[speed_at + len(b"Test_Speed"):]
            candidates = [
                struct.unpack_from("<f", tail, offset)[0]
                for offset in range(0, min(32, max(0, len(tail) - 3)))
            ]
            candidates = [v * 0.0254 for v in candidates if np.isfinite(v) and 0.001 <= v * 0.0254 <= 20]
            if candidates:
                speed = min(candidates, key=lambda value: abs(value - speed))
        runs.append((speed, displacement[:count], force[:count], velocity[:count]))
    if not runs:
        raise ValueError("PVP Session stream contains no displacement/force/velocity runs")
    return runs


def _evaluate_run(path: Path, format_name: str, current_a: float, run_index: int, speed: float, displacement, force, velocity):
    count = min(len(displacement), len(force), len(velocity))
    rate = max(100.0, count * max(speed, 0.01) / max(float(np.ptp(displacement)) / 1000.0, 1e-6) / 4.0)
    frame = pd.DataFrame({
        TIME: np.arange(count, dtype=float) / rate,
        DISP: np.asarray(displacement[:count], dtype=float),
        LOAD: np.asarray(force[:count], dtype=float),
        CURRENT: float(current_a),
        "Block ID": 1,
    })
    bounds = detect_complete_cycles(frame, AnalyzerConfig(profile=EvaluationProfile.AUDI))
    if not bounds:
        raise ValueError(f"No complete cycle at {speed:g} m/s")
    start, end = bounds[-1]
    cycle = frame.iloc[start : end + 1]
    x = cycle[DISP].to_numpy(float)
    f = cycle[LOAD].to_numpy(float)
    v = np.asarray(velocity[start : end + 1], dtype=float)
    center, stroke = float((x.max() + x.min()) / 2), float(np.ptp(x))
    window = np.abs(x - center) <= 0.05 * stroke
    rebound = f[window & (v > 0)]
    compression = f[window & (v < 0)]
    if not len(rebound) or not len(compression):
        raise ValueError(f"No directional center-window samples at {speed:g} m/s")
    return {
        "Source File": path.name,
        "Source Path": str(path),
        "Format": format_name,
        "Run": run_index,
        "Current A": current_a,
        "Speed m/s": speed,
        "Rebound N": float(np.max(rebound)),
        "Compression N": float(np.min(compression)),
        "Cycle Samples": int(end - start + 1),
        "Center Window Samples": int(np.count_nonzero(window)),
        "Stroke mm": stroke,
    }


def inspect_map_file(path: str | Path, current_a: float | None = None) -> ImportedMapFile:
    path = Path(path)
    current = current_from_filename(path) if current_a is None else float(current_a)
    suffix = path.suffix.lower()
    runs = _parse_pvp(path, current) if suffix == ".pvp" else _parse_dctw(path, current) if suffix == ".dctw" else None
    if runs is None:
        raise ValueError(f"Unsupported map file: {suffix}")
    return ImportedMapFile(path, current, suffix[1:].upper(), len(runs), tuple(float(run[0]) for run in runs))


def analyze_map_files(files: list[ImportedMapFile], max_speed_mps: float = 1.047) -> MapAnalysisResult:
    currents = {round(float(item.current_a), 9) for item in files if np.isfinite(item.current_a)}
    if len(currents) < 2:
        raise ValueError("第20/21项至少需要两个不同电流档 / Items 20/21 require at least two different current levels")
    rows, file_rows = [], []
    for item in files:
        try:
            runs = _parse_pvp(item.path, item.current_a) if item.format.upper() == "PVP" else _parse_dctw(item.path, item.current_a)
            kept = 0
            for run_index, (speed, displacement, force, velocity) in enumerate(runs, 1):
                if speed <= max_speed_mps * 1.001:
                    rows.append(_evaluate_run(item.path, item.format, item.current_a, run_index, speed, displacement, force, velocity))
                    kept += 1
            file_rows.append({"File": item.path.name, "Path": str(item.path), "Format": item.format, "Current A": item.current_a, "Run Count": len(runs), "Used Runs": kept, "Status": "OK"})
        except Exception as exc:
            file_rows.append({"File": item.path.name, "Path": str(item.path), "Format": item.format, "Current A": item.current_a, "Run Count": 0, "Used Runs": 0, "Status": str(exc)})
    detail = pd.DataFrame(rows)
    if detail.empty:
        raise ValueError("No valid damping-force runs at or below 1.047 m/s")
    long = detail.melt(
        id_vars=["Source File", "Source Path", "Format", "Run", "Current A", "Speed m/s"],
        value_vars=["Rebound N", "Compression N"], var_name="Direction", value_name="Force N",
    )
    long["Direction"] = long["Direction"].str.replace(" N", "", regex=False)
    long["Abs Force N"] = long["Force N"].abs()
    grouped = long.groupby(["Current A", "Speed m/s", "Direction"], as_index=False).agg(
        **{"Force N": ("Force N", "mean"), "Abs Force N": ("Abs Force N", "mean"), "Force SD N": ("Force N", lambda s: float(s.std(ddof=0))), "Repeat Count": ("Source File", "count")}
    )
    linearity_rows, spread_rows = [], []
    for (speed, direction), group in grouped.groupby(["Speed m/s", "Direction"], sort=True):
        group = group.sort_values("Current A")
        soft_index, hard_index = group["Abs Force N"].idxmin(), group["Abs Force N"].idxmax()
        soft, hard = float(group.loc[soft_index, "Abs Force N"]), float(group.loc[hard_index, "Abs Force N"])
        denominator = hard - soft
        normalized = (
            (group["Abs Force N"] - soft) / denominator
            if denominator > 0
            else pd.Series(np.nan, index=group.index)
        )
        if len(group) >= 2 and denominator > 0:
            coefficient = np.polyfit(group["Current A"], normalized, 1)
            fitted = np.polyval(coefficient, group["Current A"])
            residual = normalized - fitted
            ss_total = float(np.sum((normalized - normalized.mean()) ** 2))
            r_squared = 1.0 - float(np.sum(residual**2)) / ss_total if ss_total else np.nan
            maximum_deviation = float(np.max(np.abs(residual)))
        else:
            r_squared = maximum_deviation = np.nan
        for (_, row), value in zip(group.iterrows(), normalized):
            linearity_rows.append({
                "Speed m/s": speed, "Direction": direction, "Current A": row["Current A"], "Force N": row["Force N"],
                "Normalized Force": float(value) * (1 if direction == "Rebound" else -1), "Soft Force N": soft,
                "Hard Force N": hard, "Soft Current A": group.loc[soft_index, "Current A"], "Hard Current A": group.loc[hard_index, "Current A"],
                "Linearity R²": r_squared, "Max Linear Fit Deviation": maximum_deviation, "Repeat Count": row["Repeat Count"],
            })
        spread_rows.append({
            "Speed m/s": speed, "Direction": direction,
            "Damping Force Spread N": (hard - soft) * (1 if direction == "Rebound" else -1),
            "Amplification": hard / soft if soft > 0 else np.nan,
            "Soft Force N": soft, "Hard Force N": hard,
            "Soft Current A": group.loc[soft_index, "Current A"], "Hard Current A": group.loc[hard_index, "Current A"],
        })
    return MapAnalysisResult(
        pd.DataFrame(file_rows), detail, pd.DataFrame(linearity_rows), pd.DataFrame(spread_rows),
        {"Standard": "Audi VR-EF-33-2p5 sections 20/21", "Maximum Speed m/s": max_speed_mps, "Evaluation": "last complete cycle; center 10% full stroke; directional extrema", "Repeat Handling": "mean of retained files at equal current/speed/direction"},
    )


def export_map_analysis_xlsx(result: MapAnalysisResult, path: str | Path) -> Path:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    output = Path(path).with_suffix(".xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        result.current_force_linearity.to_excel(writer, sheet_name="20 Current-Force Linearity", index=False)
        result.spread_amplification.to_excel(writer, sheet_name="21 Spread-Amplification", index=False)
        result.run_detail.to_excel(writer, sheet_name="Run Detail", index=False)
        result.files.to_excel(writer, sheet_name="Source Files", index=False)
        pd.DataFrame([{"Setting": k, "Value": v} for k, v in result.settings.items()]).to_excel(writer, sheet_name="Settings", index=False)
    workbook = load_workbook(output)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for index, cells in enumerate(sheet.columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(50, max(10, max((len(str(c.value)) for c in cells if c.value is not None), default=8) + 2))
    workbook.save(output)
    return output
